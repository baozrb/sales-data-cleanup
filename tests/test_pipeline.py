"""Tests for the cleaning rules and an end-to-end run on the generated sample exports."""
import random
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import clean
import config
import generate_data
from clean import load_source, parse_dates, parse_money, process, standardize
from run import run

PRODUCTS = pd.DataFrame({"sku": ["MUG-001", "TEA-010"], "product": ["Mug", "Tea"],
                         "category": ["Kitchen", "Tea"]})


def write_csv(path, text):
    path.write_text(text, encoding="utf-8-sig")
    return path


def test_day_first_dates_and_real_excel_date_cells():
    values = pd.Series(["03/04/2026", datetime(2026, 4, 3), "31/12/2026", "not a date", None], dtype=object)
    parsed = parse_dates(values, "%d/%m/%Y")
    assert list(parsed[:3]) == [pd.Timestamp("2026-04-03"), pd.Timestamp("2026-04-03"),
                                pd.Timestamp("2026-12-31")]
    assert parsed[3:].isna().all()


def test_timestamps_are_converted_to_local_time_before_taking_the_date():
    # 02:30 UTC on 1 March is still 28 February in New York
    parsed = parse_dates(pd.Series(["2026-03-01T02:30:00+00:00"]), "ISO8601")
    assert parsed[0] == pd.Timestamp("2026-02-28")


def test_money_parsing():
    parsed = parse_money(pd.Series(["$1,234.50", " 12 ", 7.5, "USD 3.00", "", None, "abc"], dtype=object))
    assert list(parsed[:4]) == [1234.5, 12.0, 7.5, 3.0]
    assert parsed[4:].isna().all()


def test_shopify_order_fields_are_filled_down_to_every_line_item(tmp_path):
    path = write_csv(tmp_path / "shopify_orders_x.csv",
                     "Name,Email,Financial Status,Created at,Lineitem quantity,Lineitem price,Lineitem sku\n"
                     "#1001,A@Example.com,paid,2026-01-05 10:00:00 -0500,1,14.00,mug-001 \n"
                     "#1001,,,,2,9.50,TEA-010\n")
    cfg = config.SOURCES["Shopify"]
    df = standardize(load_source("Shopify", cfg, path), cfg)
    assert list(df["email"]) == ["a@example.com"] * 2
    assert list(df["status"]) == ["paid"] * 2
    assert list(df["sku"]) == ["MUG-001", "TEA-010"]
    assert list(df["source_row"]) == [2, 3]


def test_amazon_line_total_becomes_unit_price(tmp_path):
    path = tmp_path / "amazon_orders_x.txt"
    path.write_text("amazon-order-id\tpurchase-date\tsku\tquantity\titem-price\torder-status\n"
                    "111-1\t2026-01-05T15:00:00+00:00\tTEA-010\t3\t28.50\tShipped\n")
    cfg = config.SOURCES["Amazon"]
    df = standardize(load_source("Amazon", cfg, path), cfg)
    assert df["unit_price"].iloc[0] == pytest.approx(9.5)


def shopify_frame(tmp_path, name, rows):
    header = "Name,Email,Financial Status,Created at,Lineitem quantity,Lineitem price,Lineitem sku\n"
    path = write_csv(tmp_path / name, header + "".join(rows))
    cfg = config.SOURCES["Shopify"]
    return standardize(load_source("Shopify", cfg, path), cfg)


def test_newer_export_wins_and_its_refund_removes_the_order(tmp_path):
    older = shopify_frame(tmp_path, "shopify_orders_1.csv", ["#1001,a@x.com,paid,2026-03-30 10:00:00 -0400,1,14.00,MUG-001\n"])
    newer = shopify_frame(tmp_path, "shopify_orders_2.csv", ["#1001,a@x.com,refunded,2026-03-30 10:00:00 -0400,1,14.00,MUG-001\n"])
    kept, removed = process([older, newer], PRODUCTS)
    assert kept.empty
    assert list(removed["reason"]) == [clean.REPLACED, clean.CANCELLED]
    assert list(removed["source_file"]) == ["shopify_orders_1.csv", "shopify_orders_2.csv"]


def test_exact_duplicates_are_removed_once(tmp_path):
    row = "#1001,a@x.com,paid,2026-01-05 10:00:00 -0500,1,14.00,MUG-001\n"
    df = shopify_frame(tmp_path, "shopify_orders_1.csv", [row, row, row])
    kept, removed = process([df], PRODUCTS)
    assert len(kept) == 1
    assert list(removed["reason"]) == [clean.EXACT_DUP] * 2


def test_bad_rows_are_removed_with_a_reason(tmp_path):
    df = shopify_frame(tmp_path, "shopify_orders_1.csv", [
        "#1001,a@x.com,paid,2026-01-05 10:00:00 -0500,1,14.00,MUG-001\n",   # good
        "#1002,a@x.com,paid,2026-01-05 10:00:00 -0500,-1,14.00,MUG-001\n",  # return
        "#1003,a@x.com,paid,2026-01-05 10:00:00 -0500,1,,MUG-001\n",        # no price
        "#1004,a@x.com,paid,2026-01-05 10:00:00 -0500,1,14.00,MUG001\n",    # typo in SKU
        "#1005,a@x.com,paid,yesterday,1,14.00,MUG-001\n",                   # bad date
    ])
    kept, removed = process([df], PRODUCTS)
    assert list(kept["order_id"]) == ["#1001"]
    assert dict(zip(removed["order_id"], removed["reason"])) == {
        "#1002": clean.BAD_QTY, "#1003": clean.BAD_PRICE, "#1004": clean.UNKNOWN_SKU, "#1005": clean.BAD_DATE}


@pytest.fixture(scope="module")
def sample_run(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("run")
    raw, products = tmp / "raw", tmp / "products.csv"
    raw.mkdir()
    saved = config.RAW_DIR, config.PRODUCTS_FILE
    config.RAW_DIR, config.PRODUCTS_FILE = raw, products  # generate into tmp, not data/
    try:
        rng = random.Random(42)
        generate_data.products()
        generate_data.shopify(rng)
        generate_data.amazon(rng)
        generate_data.store(rng)
    finally:
        config.RAW_DIR, config.PRODUCTS_FILE = saved
    return (tmp / "out", *run(raw, tmp / "out", products))


def test_every_row_read_is_in_the_report_or_listed_with_a_reason(sample_run):
    out, kept, removed, rows_read = sample_run
    assert sum(rows_read.values()) == len(kept) + len(removed)
    assert set(removed["reason"]) <= set(clean.REASONS)
    assert not kept.duplicated(["source", "order_id", "sku"]).any()


def test_report_totals_match_the_clean_data(sample_run):
    out, kept, removed, rows_read = sample_run
    wb = load_workbook(out / "sales_report.xlsx")
    assert wb.sheetnames == ["Summary", "Products", "Clean Data", "Data Quality", "Removed Rows"]
    assert wb["Summary"]["A5"].value == pytest.approx(kept["revenue"].sum())
    by_month = pd.read_excel(out / "sales_report.xlsx", sheet_name="Summary", skiprows=7)
    assert by_month["Total"].sum() == pytest.approx(kept["revenue"].sum())
    quality = pd.read_excel(out / "sales_report.xlsx", sheet_name="Data Quality", index_col=0,
                            nrows=len(clean.REASONS) + 2)
    totals = quality["Total"].dropna()
    assert totals["Rows read"] + totals.drop(["Rows read", "Rows in report"]).sum() == totals["Rows in report"]
